import boto3
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

# Initialize Cost Explorer client
ce = boto3.client('ce')

# Date range: yesterday to today
end = datetime.today().date()
start = end - timedelta(days=1)
start_date = start.strftime('%Y-%m-%d')
end_date = end.strftime('%Y-%m-%d')

# Call AWS Cost Explorer API
response = ce.get_cost_and_usage(
    TimePeriod={
        'Start': start_date,
        'End': end_date
    },
    Granularity='DAILY',
    Metrics=['UnblendedCost'],
    GroupBy=[
        {
            'Type': 'DIMENSION',
            'Key': 'SERVICE'
        }
    ]
)

# Extract results
results = response['ResultsByTime'][0]['Groups']
total_cost = 0.0

# Create a new Excel workbook
wb = Workbook()
ws = wb.active
ws.title = f"AWS Cost {start_date}"

# Write headers
headers = ["Service", "Cost (USD)"]
ws.append(headers)

# Apply header style
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

# Write service cost data
for group in results:
    service = group['Keys'][0]
    amount = float(group['Metrics']['UnblendedCost']['Amount'])
    ws.append([service, round(amount, 2)])
    total_cost += amount

# Append total
ws.append(["Total", round(total_cost, 2)])
ws[f"A{ws.max_row}"].font = Font(bold=True)
ws[f"B{ws.max_row}"].font = Font(bold=True)

# Auto-adjust column widths
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        except:
            pass
    ws.column_dimensions[column].width = max_length + 2

# Save Excel file
filename = f"aws_cost_report_{start_date}.xlsx"
wb.save(filename)

print(f"✅ Report saved to {filename}")

